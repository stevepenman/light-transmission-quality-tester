"""Regression checks for offline boundaries, invalid readings, and Excel values."""

import json
import os
from pathlib import Path
import socket
import tempfile
import unittest
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PyQt5.QtWidgets import QApplication
from openpyxl import load_workbook

import ui
from excel_export import ExcelReportGenerator


class PortfolioRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = mock.patch.dict(os.environ, {"WTT_DATA_DIR": self.temp.name})
        self.env.start()
        self.network = mock.patch.object(socket.socket, "connect", side_effect=AssertionError("Network forbidden"))
        self.network.start()
        self.gui = ui.BottleTesterGUI()
        self.gui.show_success = mock.Mock()
        self.gui.show_warning = mock.Mock()

    def tearDown(self):
        self.gui.close()
        self.network.stop()
        self.env.stop()
        self.temp.cleanup()

    def test_every_delivery_entry_point_is_inert_even_with_settings(self):
        self.gui.email_settings.update(
            recipients=["demo@example.test"], smtp_server="smtp.example.test",
            username="demo@example.test", password="fictional-demo-value"
        )
        self.assertIsNone(self.gui.start_weekly_email_scheduler())
        self.assertFalse(self.gui.send_weekly_email())
        self.assertFalse(self.gui.send_test_email())
        self.assertFalse(self.gui.send_threshold_change_email([]))
        self.assertFalse(self.gui.send_calibration_update_email("", "", "", {}))
        self.assertFalse(self.gui.send_calibration_unlock_email("", "", "", {}, ""))

    def test_invalid_readings_are_rejected_without_recording(self):
        for values in ([float("nan"), 1, 1, 1], [float("-inf"), 1, 1, 1],
                       [-1, 1, 1, 1], [1, 1, 1], [1, 1, 1, 1, 1]):
            with self.subTest(values=values):
                self.assertEqual("ERROR", self.gui.evaluate_results(values, True, "single"))
        self.assertEqual([], self.gui.history_entries)

    def test_demo_gate_changes_do_not_write_credentials(self):
        self.gui.new_password_input.setText("fictional-session-value")
        self.gui.confirm_password_input.setText("fictional-session-value")
        self.gui.change_setup_password()
        self.assertTrue(self.gui.is_valid_password("fictional-session-value"))
        self.assertFalse(Path(self.gui.email_settings_file).exists())

    def test_report_contains_positive_ui_readings_and_synthetic_counts(self):
        examples = Path(__file__).resolve().parents[1] / "examples"
        history = json.loads((examples / "synthetic_weekly_history.json").read_text())
        calibration = json.loads((examples / "synthetic_calibration.json").read_text())
        path = ExcelReportGenerator().generate_weekly_report(history, {"total_tests_count": 4}, calibration)
        self.assertIsNotNone(path)
        self.assertTrue(Path(path).is_relative_to(Path(self.temp.name)))
        with open(path, "rb") as source:
            workbook = load_workbook(source)
        try:
            self.assertEqual(["Summary", "Detailed Data", "Statistics", "Calibration"], workbook.sheetnames)
            sheet = workbook["Detailed Data"]
            self.assertEqual((60, 65, 70, 75), tuple(sheet.cell(2, col).value for col in range(6, 10)))
            self.assertEqual("A+2", sheet["E2"].value)
            self.assertEqual(2, sum(sheet.cell(row, 4).value == "PASS" for row in range(2, 6)))
            self.assertEqual(2, sum(sheet.cell(row, 4).value == "FAIL" for row in range(2, 6)))
        finally:
            workbook.close()

    def test_empty_report_is_supported(self):
        path = ExcelReportGenerator().generate_weekly_report([], {}, {})
        self.assertIsNotNone(path)
        with open(path, "rb") as source:
            workbook = load_workbook(source)
        try:
            self.assertEqual("No data", workbook["Summary"]["B4"].value)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
